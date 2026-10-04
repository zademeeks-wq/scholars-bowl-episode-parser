#!/usr/bin/env python3
"""Merge the per-episode CSVs in data/csv into the combined sheet files.

Writes, in data/:
  Questions.csv       every row: the 19 columns of the S42 Questions tab, plus Needs Review and Review Notes
  review_queue.csv    only the rows with Needs Review = Yes, with the caption excerpt that supports each one
  Questions.xlsx      the same two tables as tabs ("Questions", "Review Queue") plus a "Notes" tab,
                      formatted like the S42 workbook (Arial 10, dark blue header row, frozen header, filter)

Usage:  python3 tools/assemble.py [--no-xlsx]
"""
import argparse
import csv
import glob
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_DIR = os.path.join(ROOT, 'data', 'csv')
OUT_DIR = os.path.join(ROOT, 'data')

S42_COLUMNS = ['Episode', 'Round', 'Question Type', 'Question #', 'Question Text', 'Correct Answer',
               'Category', 'Subject Area', 'Matchup', 'Answered By', 'Team', 'Result', 'Steal By',
               'Steal Team', 'Steal Result', 'Outcome', 'Points', 'Points To', 'Parts']
QUESTION_COLUMNS = S42_COLUMNS + ['Needs Review', 'Review Notes']
REVIEW_COLUMNS = ['Questions Row', 'Episode', 'Round', 'Question Type', 'Question #', 'Question Text',
                  'Correct Answer', 'Matchup', 'Answered By', 'Team', 'Result', 'Steal By', 'Steal Team',
                  'Steal Result', 'Outcome', 'Points', 'Points To', 'Review Notes', 'Transcript Excerpt']

# Column widths copied from the S42 Questions tab (characters); the two added columns follow.
S42_WIDTHS = [11.57, 12.0, 12.29, 12.43, 72.86, 32.0, 17.0, 20.0, 34.0, 13.0, 30.0, 10.0, 12.0, 30.0,
              12.0, 10.0, 8.0, 30.0, 7.0]
QUESTION_WIDTHS = S42_WIDTHS + [12.0, 70.0]
REVIEW_WIDTHS = [11.0, 11.57, 12.0, 12.29, 10.0, 60.0, 28.0, 34.0, 13.0, 30.0, 10.0, 12.0, 30.0, 12.0,
                 10.0, 8.0, 30.0, 70.0, 90.0]

# Numeric columns are written as numbers; everything else is text.
NUMERIC = {'Question #', 'Points', 'Parts', 'Questions Row'}

NOTES = [
    ('Notes', None),
    (None, None),
    ('Scope', None),
    ('Seasons 40 and 41 of Tennessee Scholars Bowl, built from the PBS caption files season_40.txt and '
     'season_41.txt. The Questions tab follows the 19 columns of the Season 42 Questions tab; Needs Review '
     'and Review Notes are added at the end.', None),
    ('Season 40 has episodes 1 to 10 and 17 to 62 in the caption file (episodes 11 to 16 are not in it). '
     'Season 41 has episodes 39 to 51, 53, 55 and 59 with usable text; episodes 35 to 38, 52, 54 and 57 are '
     'empty, and episodes 56 and 58 repeat the text of episode 59, so only episode 59 is recorded.', None),
    (None, None),
    ('How rows are filled', None),
    ('Tossup 10 points; bonus 20 points unless the host states another value (for example 15 points for '
     'three of four answers, 5 per item on "name four" bonuses). A two-part bonus uses Part 1 / Part 2 in the '
     'Question Text and Correct Answer and has Parts = 2; a half-right answer is Partial.', None),
    ('Answered By is the first player to buzz. Steal By is the player on the other team who answered after a '
     'wrong first answer. Team, Steal Team and Points To use the school name; a school with more than one '
     'team in a season is labelled with its captain, for example Maryville High School (Mohammed\'s team).', None),
    ('When the answer is a halftime challenge or the host re-ruled it, the final ruling is recorded and the '
     'row has a note.', None),
    ('The captions have no speaker labels and many names are garbled. Player names are first names read '
     'from the host\'s introductions and may be misspelled.', None),
    (None, None),
    ('Review columns', None),
    ('Needs Review = Yes marks a row whose buzzer, team, answer, question text or ruling could not be read '
     'reliably from the captions. Review Notes says what is uncertain. Team = Unknown means the captions do '
     'not name the buzzer or the team that took the question; those rows also have Points To = Unknown. '
     'A Result of Unclear has no points counted.', None),
    ('The Review Queue tab lists the flagged rows with the caption lines they came from, so each can be '
     'checked without going back to the source files.', None),
    (None, None),
    ('Counts', None),
]
NOTES_FORMULAS = [
    ('Rows on the Questions tab', '=COUNTA(Questions!A2:A)'),
    ('Tossups', '=COUNTIF(Questions!C2:C,"Tossup")'),
    ('Bonuses', '=COUNTIF(Questions!C2:C,"Bonus")'),
    ('Rows marked Needs Review', '=COUNTIF(Questions!T2:T,"Yes")'),
    ('Rows with Team = Unknown', '=COUNTIF(Questions!K2:K,"Unknown")'),
]


def episode_key(name):
    m = re.match(r'S(\d+)_E(\d+)', name)
    return (int(m.group(1)), int(m.group(2)))


def load_rows():
    rows = []
    for path in sorted(glob.glob(os.path.join(CSV_DIR, 'S*_E*.csv')),
                       key=lambda p: episode_key(os.path.basename(p))):
        with open(path, newline='', encoding='utf-8') as fh:
            rows.extend(csv.DictReader(fh))
    return rows


def check(rows):
    problems = []
    seen = set()
    for r in rows:
        key = (r['Episode'], r['Question Type'], r['Question #'])
        if key in seen:
            problems.append('duplicate row %s' % (key,))
        seen.add(key)
        for col in ('Question Text', 'Correct Answer', 'Review Notes'):
            if '—' in r.get(col, ''):
                problems.append('em dash in %s %s %s' % (r['Episode'], r['Question Type'], r['Question #']))
        if r['Result'] == 'Unclear' and r['Needs Review'] != 'Yes':
            problems.append('unclear but not flagged: %s' % (key,))
    return problems


def as_number(value):
    if value in ('', None):
        return None
    return int(float(value))


def table(rows, columns):
    out = []
    for r in rows:
        out.append([as_number(r[c]) if c in NUMERIC else r[c] for c in columns])
    return out


def write_csv(path, columns, rows):
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        w.writerow(columns)
        for r in rows:
            w.writerow(['' if v is None else v for v in r])


def write_xlsx(path, questions, review):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F4E78')
    body_font = Font(name='Arial', size=10)

    def sheet(ws, columns, widths, rows, wrap_cols):
        ws.append(columns)
        for row in rows:
            ws.append(row)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical='top', wrap_text=columns[cell.column - 1] in wrap_cols)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = 'Questions'
    sheet(ws, QUESTION_COLUMNS, QUESTION_WIDTHS, questions,
          {'Question Text', 'Correct Answer', 'Matchup', 'Team', 'Steal Team', 'Points To', 'Review Notes'})
    last = ws.max_row
    ws.conditional_formatting.add(
        'A2:U%d' % last,
        FormulaRule(formula=['$T2="Yes"'], fill=PatternFill('solid', bgColor='FFF2CC', fgColor='FFF2CC')))

    ws2 = wb.create_sheet('Review Queue')
    sheet(ws2, REVIEW_COLUMNS, REVIEW_WIDTHS, review, {'Question Text', 'Review Notes', 'Transcript Excerpt'})

    ws3 = wb.create_sheet('Notes')
    ws3.column_dimensions['A'].width = 110
    ws3.column_dimensions['B'].width = 12
    for text, _ in NOTES:
        ws3.append([text])
    for label, formula in NOTES_FORMULAS:
        ws3.append([label, formula])
    for row in ws3.iter_rows():
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column == 1)
    for cell in ws3['A']:
        if cell.value in ('Notes', 'Scope', 'How rows are filled', 'Review columns', 'Counts'):
            cell.font = Font(name='Arial', size=10, bold=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--no-xlsx', action='store_true', help='skip the .xlsx file')
    args = ap.parse_args()

    rows = load_rows()
    problems = check(rows)
    for p in problems:
        print('problem:', p, file=sys.stderr)

    questions = table(rows, QUESTION_COLUMNS)
    review = []
    for i, r in enumerate(rows):
        if r['Needs Review'] == 'Yes':
            r = dict(r, **{'Questions Row': str(i + 2)})
            review.append([as_number(r[c]) if c in NUMERIC else r[c] for c in REVIEW_COLUMNS])

    write_csv(os.path.join(OUT_DIR, 'Questions.csv'), QUESTION_COLUMNS, questions)
    write_csv(os.path.join(OUT_DIR, 'review_queue.csv'), REVIEW_COLUMNS, review)
    if not args.no_xlsx:
        write_xlsx(os.path.join(OUT_DIR, 'Questions.xlsx'), questions, review)

    episodes = sorted({r['Episode'] for r in rows}, key=lambda e: (int(e[1:3]), int(e.split('E')[1])))
    print('%d rows from %d episodes (%d tossups, %d bonuses); %d flagged for review; %d problems'
          % (len(rows), len(episodes), sum(r['Question Type'] == 'Tossup' for r in rows),
             sum(r['Question Type'] == 'Bonus' for r in rows), len(review), len(problems)))
    return 1 if problems else 0


if __name__ == '__main__':
    sys.exit(main())
